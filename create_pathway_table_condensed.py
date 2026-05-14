#!/usr/bin/env python3
"""
Condensed Table 4 — only the key pathways highlighted in the results text.
Same style as Table4_Pathways_preterm_infants.docx (template-matched).
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "pathway_results/Table4_Pathways_preterm_infants_condensed.docx"
FS = 9
FS_CAP = 10

# ── Helpers (identical to main script) ───────────────────────────────────────

def fmt_p(v):
    if v < 0.001:  return "<0.001"
    if v < 0.01:   return f"{v:.4f}"
    return f"{v:.3f}"

def add_run(para, text, bold=False, italic=False, size=None):
    r = para.add_run(text)
    r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return r

def _border_xml(name, style):
    el = OxmlElement(f"w:{name}")
    if style == "single":
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0");    el.set(qn("w:color"), "000000")
    else:
        el.set(qn("w:val"), "none")
    return el

def set_cell_borders(cell, top="none", bottom="none"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")): tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    for name, style in [("top", top), ("bottom", bottom),
                         ("left", "none"), ("right", "none"),
                         ("insideH", "none"), ("insideV", "none")]:
        b.append(_border_xml(name, style))
    tcPr.append(b)

def set_row_borders(row, top=False, bottom=False):
    for cell in row.cells:
        set_cell_borders(cell,
                         top="single" if top else "none",
                         bottom="single" if bottom else "none")

def clear_table_borders(table):
    tbl = table._tbl; tblPr = tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "none")
        tblBorders.append(el)
    tblPr.append(tblBorders)

def fill_cell(cell, text, bold=False, italic=False, size=FS):
    add_run(cell.paragraphs[0], text, bold=bold, italic=italic, size=size)

# ── Condensed data (key pathways only, per results narrative) ─────────────────

HEADERS = ["Metabolites", "Pathway Name", "Total", "Hits", "Raw P", "FDR", "Function"]
COL_WIDTHS = [Cm(3.5), Cm(4.0), Cm(1.0), Cm(1.0), Cm(1.7), Cm(1.7), Cm(3.1)]

SECTIONS = [
    {
        "label": "6M GA only",
        "note": "All GA-significant metabolites at 6M were also significant in BPD group",
        "rows": [],
    },
    {
        "label": "6M BPD only",
        "note": None,
        "rows": [
            # Both reach FDR significance
            ("N-Acetylaspartic acid, Glutamine, Citric acid, Succinic acid",
             "Alanine, aspartate and glutamate metabolism",
             28, 4, 9.1045e-06, 0.00072836, "Amino acid metabolism"),
            ("Acetic acid, Glutamine, Citric acid",
             "Glyoxylate and dicarboxylate metabolism",
             32, 3, 0.00057151, 0.022861, "Carbohydrate metabolism"),
        ],
    },
    {
        "label": "6M Both GA+BPD",
        "note": None,
        "rows": [
            # Key pathway reflecting the 3-MOV signal across both groupings
            ("3-Methyl-2-oxopentanoic acid, Valine",
             "Valine, leucine and isoleucine biosynthesis",
             8, 2, 0.00078202, 0.062561, "Amino acid metabolism"),
        ],
    },
    {
        "label": "2Y GA only",
        "note": None,
        "rows": [
            # GA-specific findings; neither reaches FDR significance
            ("Tyrosine, 4-Hydroxyphenylacetate",
             "Tyrosine metabolism",
             42, 2, 0.017206, 0.79173, "Amino acid metabolism"),
            ("Taurine",
             "Taurine and hypotaurine metabolism",
             8, 1, 0.039587, 0.79173, "Metabolism of other amino acids"),
        ],
    },
    {
        "label": "2Y BPD only",
        "note": "Too few metabolites to perform meaningful enrichment analysis "
                "(2-Aminobutyric acid, Citric acid, Propylene glycol)",
        "rows": [],
    },
    {
        "label": "2Y Both GA+BPD",
        "note": None,
        "rows": [
            # Persistent amino acid and carbohydrate metabolic disruption
            ("Glutamine, Succinic acid",
             "Alanine, aspartate and glutamate metabolism",
             28, 2, 0.0042853, 0.22341, "Amino acid metabolism"),
            ("Glycine, Glutamine",
             "Glyoxylate and dicarboxylate metabolism",
             32, 2, 0.0055852, 0.22341, "Carbohydrate metabolism"),
        ],
    },
]

# ── Build document ────────────────────────────────────────────────────────────

def main():
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.54)
        sec.top_margin  = sec.bottom_margin = Cm(2.54)

    # Caption
    cap = doc.add_paragraph()
    add_run(cap, "Table 4.", bold=True, size=FS_CAP)
    add_run(cap, (
        " Key metabolic pathways enriched among urinary NMR metabolites "
        "significantly expressed in preterm infants stratified by gestational "
        "age (GA) and bronchopulmonary dysplasia (BPD) severity at 6 months "
        "(6M) and 2 years (2Y) corrected age."
    ), size=FS_CAP)

    # Count rows
    n_rows = 1
    for s in SECTIONS:
        n_rows += 1
        if s["note"] and not s["rows"]: n_rows += 1
        n_rows += len(s["rows"])

    table = doc.add_table(rows=n_rows, cols=7)
    clear_table_borders(table)
    for row in table.rows:
        for i, w in enumerate(COL_WIDTHS):
            row.cells[i].width = w

    # Header
    hdr = table.rows[0]
    for i, h in enumerate(HEADERS):
        p = hdr.cells[i].paragraphs[0]
        if h == "Raw P":
            add_run(p, "Raw ", bold=True, size=FS)
            add_run(p, "P", bold=True, italic=True, size=FS)
        else:
            add_run(p, h, bold=True, size=FS)
    set_row_borders(hdr, top=True, bottom=True)

    # Sections
    ri = 1
    for si, sec in enumerate(SECTIONS):
        is_last_sec = si == len(SECTIONS) - 1

        sec_row = table.rows[ri]
        fill_cell(sec_row.cells[0], sec["label"], bold=True)
        set_row_borders(sec_row, top=True, bottom=False)
        ri += 1

        if sec["note"] and not sec["rows"]:
            note_row = table.rows[ri]
            merged = note_row.cells[0].merge(note_row.cells[6])
            fill_cell(merged, sec["note"], italic=True)
            set_cell_borders(merged, top="none", bottom="none")
            ri += 1

        for di, row_data in enumerate(sec["rows"]):
            met, pwy, total, hits, rawp, fdr, func = row_data
            is_last_row = is_last_sec and di == len(sec["rows"]) - 1
            dr = table.rows[ri]
            for i, v in enumerate([met, pwy, str(total), str(hits),
                                    fmt_p(rawp), fmt_p(fdr), func]):
                fill_cell(dr.cells[i], v)
            set_row_borders(dr, top=False, bottom=is_last_row)
            ri += 1

    # Footnote
    fn = doc.add_paragraph()
    add_run(fn, "Total, total number of compounds in the pathway; "
                "Hits, matched number from the uploaded data; Raw ", size=FS_CAP)
    add_run(fn, "P", italic=True, size=FS_CAP)
    add_run(fn, ", original ", size=FS_CAP)
    add_run(fn, "P", italic=True, size=FS_CAP)
    add_run(fn, (" value from the enrichment analysis; FDR, false discovery rate "
                 "(Benjamini-Hochberg correction)."), size=FS_CAP)

    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")

if __name__ == "__main__":
    main()
